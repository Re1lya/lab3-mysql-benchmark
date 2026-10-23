from pathlib import Path
import json,copy,hashlib,csv,math
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
REPO=Path(__file__).resolve().parents[2]
ROOT=REPO/'report/assets'
OUT=REPO/'report';OUT.mkdir(parents=True,exist_ok=True)
REF=OUT/'templates/paper-template.docx'
D=Document(REF)
for child in list(D._element.body):
 if child.tag!=qn('w:sectPr'):D._element.body.remove(child)
# Preserve reference geometry, replacing source-specific text and typography fallbacks.
for s in D.styles:
 if s.type==1:
  s.font.color.rgb=RGBColor(0,0,0)
  rf=s.element.get_or_add_rPr().find(qn('w:rFonts'))
  if rf is None:rf=OxmlElement('w:rFonts');s.element.get_or_add_rPr().append(rf)
  for attr in list(rf.attrib):
   if 'Theme' in attr:del rf.attrib[attr]
  rf.set(qn('w:eastAsia'),'黑体' if s.name in ['Title','Heading 1','Heading 2','Heading 3'] else '宋体')
  rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')
  for w in list(s.element.get_or_add_rPr().findall(qn('w:w'))):s.element.get_or_add_rPr().remove(w)
for name in ['Normal','Body Text']:
 st=D.styles[name];st.font.size=Pt(10.5);st.paragraph_format.line_spacing=Pt(14);st.paragraph_format.space_after=Pt(4);st.paragraph_format.first_line_indent=Pt(21)
 st.paragraph_format.keep_with_next=False
for name in ['Heading 1','Heading 2','Heading 3']:
 st=D.styles[name];st.font.size=Pt(10.5);st.font.bold=True;st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True;st.paragraph_format.first_line_indent=Pt(0)
D.styles['Title'].font.size=Pt(18);D.styles['Title'].font.bold=True
for sect in D.sections:
 for hdr in [sect.header,sect.first_page_header,sect.even_page_header]:
  for p in hdr.paragraphs:p.clear()
  p=hdr.paragraphs[0];p.text='JavaEE 实验报告    MySQL读写性能测试';p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  for r in p.runs:r.font.size=Pt(8);r.font.name='宋体';r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
 for f in [sect.footer,sect.first_page_footer,sect.even_page_footer]:
  for p in f.paragraphs:p.clear()
  p=f.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  run=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
D.core_properties.author='温家正 姚景云 杜奕含 陈浩贤';D.core_properties.title='MySQL读写性能测试实验报告';D.core_properties.subject='Spring Data JPA与MyBatis读写测试及三轮复核';D.core_properties.last_modified_by=''
R=json.loads((REPO/'results/实验数据/三轮复核/统计结果.json').read_text(encoding='utf-8'))
def get(fw,op,t,run=1):return next(r for r in R if r['framework']==fw and r['operation']==op and r['threads']==t and r['run']==run)
def para(text,style='Body Text',size=None,indent=True,bold=False):
 p=D.add_paragraph(text,style);p.paragraph_format.keep_with_next=False;p.paragraph_format.widow_control=True;p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(14)
 if not indent:p.paragraph_format.first_line_indent=Pt(0)
 for r in p.runs:
  if size:r.font.size=Pt(size)
  if bold:r.bold=True
 return p
def heading(text,level=1):
 p=D.add_paragraph(text,'Heading '+str(level));return p
def page():D.add_page_break()
def table(headers,rows,widths=None):
 t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=widths or [15.5/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 pr=t._tbl.tblPr
 tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(sum(widths)*567)))
 borders=OxmlElement('w:tblBorders')
 for side in ['top','bottom','left','right','insideH','insideV']:
  x=OxmlElement('w:'+side);x.set(qn('w:val'),'single' if side in ['top','bottom'] else 'nil');x.set(qn('w:sz'),'8');x.set(qn('w:color'),'000000');borders.append(x)
 pr.append(borders)
 for i,vals in enumerate([headers]+rows):
  row=t.rows[0] if i==0 else t.add_row()
  for c,w,txt in zip(row.cells,widths,vals):
   c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.paragraph_format.left_indent=Pt(0);p.paragraph_format.right_indent=Pt(0);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   rr=p.add_run(str(txt));rr.font.size=Pt(9);rr.font.name='Times New Roman';rr._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体');rr.bold=i==0
   if i==0:
    b=OxmlElement('w:tcBorders');bottom=OxmlElement('w:bottom');bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'5');b.append(bottom);c._tc.get_or_add_tcPr().append(b)
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
 bp=D.add_paragraph();bp.paragraph_format.space_after=Pt(2);bp.paragraph_format.line_spacing=Pt(2);bp.paragraph_format.space_before=Pt(0)
 return t
def caption(txt,keep=False):
 p=para(txt,size=9,indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=keep
 return p
def picture(path,cap,width=15.1):
 p=D.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run().add_picture(str(ROOT/path),width=Cm(width));caption(cap)
# Opening mirrors the reference paper layout.
para('MySQL读写性能测试实验报告','Title',indent=False)
para('温家正  姚景云  杜奕含  陈浩贤','作者',size=12,indent=False)
para('摘  要：本实验以 Spring 商品管理应用为对象，使用 JMeter 测试查询与新建接口，对比 Spring Data JPA 与 MyBatis，并结合 Prometheus 观察运行状态。读取设置 1000～1200 线程、2 次循环，写入设置 1000～1300 线程、1 次循环，升压均为 10 s。第一轮 18 组、31200 次请求均通过既定断言。三轮原始记录共 54 组、93600 次请求，未出现断言失败，但部分轮次有显著延迟尖峰。1300 线程写入退化未稳定复现，读取上限也未测定，因此不能仅凭第一轮判断系统容量或持久层性能优劣。','摘要',size=9,indent=False)
para('关键词：MySQL；Spring Data JPA；MyBatis；JMeter；读写性能','关键词',size=9,indent=False)
para('Read and Write Performance Testing of a MySQL Based Spring Application','Title',size=12,indent=False)
para('WEN Jiazheng, YAO Jingyun, DU Yihan, CHEN Haoxian','Name',size=10,indent=False)
para('Abstract: JMeter was used to evaluate product query and creation APIs implemented with Spring Data JPA and MyBatis. The tests used a 10-second ramp-up, 1,000–1,200 threads with two loops for reads, and 1,000–1,300 threads with one loop for writes. All 31,200 requests in the first round passed the configured assertions. Review of three rounds covered 54 runs and 93,600 requests with no assertion failures, but substantial latency spikes occurred in some runs. The write slowdown at 1,300 threads was not reproduced consistently, and the read capacity limit was not established. The results therefore do not establish a fixed capacity boundary or a general performance ranking between the two persistence implementations.','Abstract',size=9,indent=False)
para('Key words: MySQL; Spring Data JPA; MyBatis; JMeter; read and write performance',size=9,indent=False)
heading('1 问题描述与实验目的')
para('数据库驱动的 Web 请求包含 HTTP 传输、业务处理、数据库访问及响应序列化。读写测试测量的是整条接口链路的表现，不能将 HTTP 响应时间直接视为 MySQL 的纯 SQL 执行时间。')
para('本实验比较商品查询与新建接口在不同负载档位下的响应时间、吞吐量和成功情况，并观察 JPA 与 MyBatis 两种实现的差异。实验同时检查重复运行的稳定性，区分已观测到的延迟现象与尚未证实的系统瓶颈。')
page()
heading('2 实验环境与系统实现')
caption('表1  服务器配置与职责',True)
table(['服务器','配置','职责与软件'],[['CDWY3','2核 / 2 GiB','编译与压测；JDK 21、Maven 3.9.16、JMeter 5.6.3'],['CDWY1','2核 / 2 GiB','应用；Spring Boot 4.1.1、JPA、MyBatis 3.5.19'],['CDWY2','2核 / 2 GiB','数据库；MySQL 9.2.0'],['CDWY4','2核 / 2 GiB','监控；Prometheus 3.7.3']],[2,3,10.5])
para('四台主机均为 Ubuntu 18.04.6，通过私网通信。应用采用 Docker Swarm 部署，各节点运行 node-exporter。实验记录中的初始数据为商品 3912 行、关联 30660 行、上架 3914 行，第一轮各组的监控元数据记录了相同初始基线。')
para('第一轮实验记录采用 Tomcat 最大/最小线程 300/100，Druid 最大/最小/初始连接 300/100/100。应用堆为 256～512 MiB；JMeter 堆为 256～768 MiB、线程栈 512 KiB。JPA 与 MyBatis 通过 profile 切换，正式统计使用相同接口与请求计划。这里记录的是 2026 年 10 月 18 日的实验条件，不以之后的服务器配置替代历史配置。')
heading('3 实验设计与测试方法')
heading('3.1 测试接口与负载设置',2)
para('读接口为 GET /shops/{shopId}/products/{id}，商品标识由 CSV 提供；写接口为 POST /shops/1/products，字段按 JMX 生成。读取检查 HTTP 200 且 errno=0，写入计划检查 HTTP 201。“断言成功”指上述校验通过，不能替代写入结果的独立数据库核对。')
caption('表2  每轮测试矩阵',True)
table(['操作','配置线程数','Ramp-up','loop','每实现请求数'],[['读取','1000 / 1050 / 1100 / 1150 / 1200','10 s','2','11000'],['写入','1000 / 1100 / 1200 / 1300','10 s','1','4600']],[1.4,6.1,2,1.2,4.8])
para('两种实现每轮共 18 组、31200 次请求。正文逐档图片与统计保持第一轮口径；第 4.4 节另列同批第二、三轮复核结果，不将不同轮次拼接为同一次测试。')
heading('3.2 指标口径与评价条件',2)
para('响应时间使用 JTL 的 elapsed，包含完整采样耗时，不扣除连接时间。吞吐量按请求数除以首个请求开始至最后请求结束的时长计算；P95、P99 用排序样本的线性插值计算，与部分 JMeter 版本的百分位算法可能有微小舍入差异。')
para('沿用参考样例的比较条件：成功且 elapsed≤40 ms 的请求占比不低于 80%。40 ms 和 80% 是本实验选取的比较标准，并非实验指导书规定的统一验收阈值；P95≤40 ms 对应另一种更严格的标准，不能混用。')
para('JMeter 在非 GUI 模式运行，每组保留 samples.jtl、jmeter.log、console.log 和起止时间。图中的 Response Times Over Time 为 1 s 聚合趋势；Response Time Percentiles 展示请求耗时分布。前者平稳并不代表每个请求都没有长尾。')
heading('4 实验结果与分析')
heading('4.1 MyBatis 查询接口测试',2)
para('本节依次展示第一轮的五个读取档位。配置 x-y-z 表示线程数、升压秒数和每线程循环数；所有图均为本组第一轮原始 JMeter 页面图像。')
readtext={1000:'该档平均响应为 3.67 ms，P95 为 6.00 ms，98.60% 的请求成功且在 40 ms 内完成。趋势图初段较高，随后下降；百分位图仍有少量超过 40 ms 的长尾，因此不能写成所有请求均在 40 ms 内。',1050:'该档平均响应为 2.01 ms，P95 为 3.00 ms，全部请求在 40 ms 内完成。与 1000 线程相比并未出现持续恶化，这说明相邻档位的均值也受运行状态影响，不能要求响应时间随线程数严格单调增加。',1100:'该档平均响应为 2.34 ms，P95 为 4.00 ms，全部请求在 40 ms 内完成。在本轮观测范围内，该档满足选定比较条件，但不能由一轮数据认定具有相同规模的持续并发承载能力。',1150:'该档平均响应为 2.15 ms，P95 为 3.00 ms，99.61% 的请求在 40 ms 内完成。少量尾部请求超过比较阈值，但总体比例仍高于 80%。',1200:'该档平均响应为 2.56 ms，P95 为 4.00 ms，99.54% 的请求在 40 ms 内完成。最高测试档仍满足比较条件，说明本轮尚未探测到读取侧的退化边界；1200 不能写为最大支持线程数。'}
fig=1
for idx,t in enumerate([1000,1050,1100,1150,1200]):
 if idx:D.add_paragraph().paragraph_format.space_after=Pt(0)
 heading(f'4.1.{idx+1} 查询测试 {t} 线程',3)
 r=get('mybatis','ReadProduct',t)
 caption(f'表{3+idx}  MyBatis 读取 {t}-10-2 第一轮',True)
 table(['线程','Ramp-up','loop','请求数','均值/ms','P95/ms','吞吐量/s','错误率'],[[t,10,2,r['n'],f"{r['mean']:.2f}",f"{r['p95']:.2f}",f"{r['rps']:.2f}",'0%']])
 picture(f'original-p{idx+3}-img2.png',f'图{fig}  {t}-10-2 Response Time Percentiles');fig+=1
 picture(f'original-p{idx+3}-img1.png',f'图{fig}  {t}-10-2 Response Times Over Time');fig+=1
 para(readtext[t])
heading('4.2 MyBatis 新建接口测试',2)
para('本节展示第一轮的四个写入档位，循环次数均为 1。请求通过 HTTP 201 断言与响应时间达到比较条件是两件不同的事。')
writetext={1000:'平均响应为 6.16 ms，P95 为 7.00 ms，98.90% 的请求在 40 ms 内成功完成。相比同线程配置的读取，平均耗时增加；但读写循环数不同，应结合吞吐量和请求分布解释差异。',1100:'平均响应为 7.30 ms，P95 为 10.00 ms，98.64% 的请求在 40 ms 内成功完成。该档整体满足比较条件，百分位图可用于识别平均值之外的少量慢请求。',1200:'平均响应为 6.18 ms，P95 为 7.00 ms，99.08% 的请求在 40 ms 内成功完成。本轮未出现相对于 1100 线程的持续退化，不能仅根据档位高低推断瓶颈。',1300:'平均响应升至 102.03 ms，P95 为 615.30 ms，40 ms 内成功比例降至 71.77%，本轮未达到 80% 比较条件。该现象在第二、三轮没有稳定复现，因此只记录为该轮延迟尖峰，不据此认定写入能力的固定上限。'}
for idx,t in enumerate([1000,1100,1200,1300]):
 if idx:D.add_paragraph().paragraph_format.space_after=Pt(0)
 heading(f'4.2.{idx+1} 新建测试 {t} 线程',3)
 r=get('mybatis','WriteProduct',t)
 caption(f'表{8+idx}  MyBatis 写入 {t}-10-1 第一轮',True)
 table(['线程','Ramp-up','loop','请求数','均值/ms','P95/ms','吞吐量/s','错误率'],[[t,10,1,r['n'],f"{r['mean']:.2f}",f"{r['p95']:.2f}",f"{r['rps']:.2f}",'0%']])
 picture(f'original-p{idx+8}-img2.png',f'图{fig}  {t}-10-1 Response Time Percentiles');fig+=1
 picture(f'original-p{idx+8}-img1.png',f'图{fig}  {t}-10-1 Response Times Over Time');fig+=1
 para(writetext[t])
heading('4.3 JPA 与 MyBatis 第一轮比较',2)
caption('表12  第一轮两种实现的响应时间与达标比例',True)
rows=[]
for op,threads in [('ReadProduct',[1000,1050,1100,1150,1200]),('WriteProduct',[1000,1100,1200,1300])]:
 for t in threads:
  a=get('jpa',op,t);b=get('mybatis',op,t)
  rows.append(['读' if op=='ReadProduct' else '写',t,f"{a['mean']:.2f}",f"{b['mean']:.2f}",f"{a['p95']:.2f}",f"{b['p95']:.2f}",f"{a['ok40']:.2f}%",f"{b['ok40']:.2f}%"])
table(['操作','线程','JPA\n均值/ms','MyBatis\n均值/ms','JPA\nP95/ms','MyBatis\nP95/ms','JPA\n≤40ms','MyBatis\n≤40ms'],rows,[1.3,1.3,2.1,2.1,2.1,2.1,2.1,2.1])
para('第一轮九个配置中，MyBatis 平均响应均低于 JPA。两种实现的十个读取组和六个写入组满足 80% 比较条件，两个 1300 线程写入组未满足。JPA 1300 线程写入均值为 3232.80 ms，说明该轮存在显著等待，但现有结果不能唯一归因于数据库锁、连接池或垃圾回收。')
para('读取与写入共有的 1000、1100、1200 线程档位中，MyBatis 读取均值低于写入。由于读取循环 2 次、写入循环 1 次，二者请求总数和到达过程不同，这一比较描述的是当前负载方案下的接口表现，并非严格等吞吐条件下的单次数据库操作成本。')
heading('4.4 三轮重复记录的稳定性复核',2)
para('2026 年 10 月 23 日从压测服务器读取 supplement-20261018 原始目录，补充核查同批次 r2、r3。共 54 组、93600 条样本，success=false 为 0。所有组均保留；第一轮图表仍对应 r1。后两轮的日志与 JTL 未包含每组完整环境和基线快照，因此该对照用于观察波动，不能单独证明环境状态完全一致。')
para('三轮中共有 49 组达到 80% 比较条件，5 组未达到。该结果反映了“功能请求完成”和“响应时间稳定”之间的区别。下面逐档列出各轮均值和达标比例，检查第一轮现象能否重复。')
page()
for op,ts,title,num in [('ReadProduct',[1000,1050,1100,1150,1200],'读取',13),('WriteProduct',[1000,1100,1200,1300],'写入',14)]:
 caption(f'表{num}  三轮{title}均值与 40 ms 内成功比例',True)
 rows=[]
 for fw in ['jpa','mybatis']:
  for t in ts:
   rs=[get(fw,op,t,i) for i in [1,2,3]]
   rows.append(['JPA' if fw=='jpa' else 'MyBatis',t]+[f"{x['mean']:.2f}" for x in rs]+[f"{x['ok40']:.2f}" for x in rs])
 table(['实现','线程','r1/ms','r2/ms','r3/ms','r1/%','r2/%','r3/%'],rows,[1.8,1.3,2.3,2.3,2.3,1.7,1.7,1.7])
para('MyBatis 1300 线程写入的三轮均值分别为 102.03、6.76、6.05 ms；JPA 对应为 3232.80、13.30、8.12 ms。第一轮的突增未重复出现。另有 JPA 1100 线程写入第二轮和 1000 线程写入第三轮出现尖峰，说明异常并不只发生在最高档。')
para('JPA 1100 线程读取第三轮的 40 ms 内成功比例为 58.05%，而同框架更高的 1150、1200 档三轮均达到比较条件。因此当前资料支持“存在轮次波动”，尚不足以支持随负载升高形成稳定拐点，也不足以宣称 MyBatis 在所有条件下必然优于 JPA。')
page();heading('4.5 Prometheus 运行监控',2)
para('以下保留第一轮 MyBatis 1300-10-1 写入的历史监控图。对应测试在北京时间 2026 年 10 月 18 日 13:13:51～13:14:04 执行；截图按原测试时间于 10 月 21 日补截，页面横轴为 UTC，需加 8 小时。')
picture('Im850.png',f'图{fig}  JVM 活跃线程数 spring:jvm_threads_live_threads');fig+=1
picture('Im855.png',f'图{fig}  应用节点 CPU 使用率 node:cpu_usage_percent');fig+=1
para('图中 JVM 活跃线程数约由 115 增至 124，应用节点 CPU 规则值约由 11.3% 增至 13.3%。JVM 线程指标统计应用进程中的线程，并不是 JMeter 配置线程数，也不等于数据库正在执行的请求数。')
heading('4.5 Prometheus 运行监控 续',2)
picture('Im867.png',f'图{fig}  应用节点内存使用率 node:memory_usage_percent');fig+=1
para('截图中应用节点内存规则值约由 35.5% 升至 38.8%。该规则为 [1−(MemFree+Buffers+Cached)/MemTotal]×100%，并非直接使用 MemAvailable 计算。')
para('本批 Prometheus 抓取间隔为 2 s，记录规则每 30 s 更新一次；CPU 规则采用 5 min rate 窗口，而单组测试仅约十余秒。因而图上的阶梯和低 CPU 均值不能排除短时资源争用，也不能证明数据库、连接池或应用线程没有瓶颈。延迟归因仍需要与原始计数器、连接池等待、GC、数据库锁及磁盘 I/O 记录对齐。')
heading('5 讨论与总结')
heading('5.1 配置线程数与负载解释',2)
para('5000-10-1 表示计划在 10 s 内依次启动 5000 个线程，每个线程执行一次测试流程；它不表示 5000 个请求同时在途。线程在请求结束后即可退出，实际活跃数量取决于启动节奏和请求耗时。JMeter 官方说明也要求结合 Ramp-up 与线程持续时间解释负载。[3]')
para('本组第一轮 MyBatis 读取配置 1000～1200 线程，JTL 样本中 allThreads 峰值实际为 22～31，吞吐量约 203～254 次/s。配置线程数、采样时活跃线程数和同时在途请求数是不同指标。另一组采用更高线程档位并不能直接说明本组数据错误，比较前必须统一升压、循环、请求内容、软件版本和资源状态。')
heading('5.2 实验结论与改进方向',2)
para('本实验完成了商品查询与新建两个接口的性能采集，保留了原始 JTL、JMeter 图表和三项运行监控，并在相同负载配置下比较 JPA 与 MyBatis。第一轮 18 组数据能够复算，三轮 54 组全部通过既定断言，但存在明显的响应时间波动。')
para('当前可得的结论是：在本批第一轮低延迟记录中，读接口通常比写接口响应更快；部分轮次会出现显著的尾部延迟，且这种退化没有形成稳定的线程档位边界。读取性能上限尚未测定，1300 线程写入也不是已证实的固定容量上限。')
para('进一步分析应采用一致且有记录的预热、数据恢复和冷却过程，延长稳态采集时间；从已验证的负载逐步增加，在出现退化的区域进行细分并重复。应同时保存压测机自身的资源与活跃线程指标，确认压力能按计划发出，再结合应用、连接池与 MySQL 监控解释延迟变化。对异常轮次保留原始证据，不能仅删去高延迟结果。')
heading('参考资料')
for t in ['[1] 实验三：MySQL读写性能测试，课程实验指导书，第1～4页。','[2] 实验报告样例，基于MySQL的Spring应用的读写效率，第1～10页。','[3] Apache Software Foundation. Apache JMeter User’s Manual: Elements of a Test Plan, Thread Group. https://jmeter.apache.org/usermanual/test_plan.html （访问日期：2026-10-23）。','[4] 本组实验程序与数据仓库：https://github.com/Re1lya/lab3-mysql-benchmark。','[5] 本组三轮原始记录：CDWY3 /root/lab3-results/supplement-20261018；实验日期2026-10-18，复核日期2026-10-23。']:
 para(t,size=9,indent=False)
heading('附录 数据与统计对应关系')
para('仓库 results/实验数据/原始数据/01_JPA读取 至 04_MyBatis写入 保留第一轮材料；新增 results/实验数据/三轮复核 保存本批三轮 JTL、日志、起止时间、逐组统计、文件校验清单与复算脚本。第一轮图表继续使用原始报告内嵌图像，没有替换为另一组数据，也没有将新绘图标称为 JMeter 截图。',size=9)
para('目录名称采用 实现-r轮次-操作-t线程数-l循环数，例如 mybatis-r1-WriteProduct-t1300-l1。每条统计记录可由对应 samples.jtl 复算。保留原文件的相对结构和 SHA-256，便于核验来源。',size=9)
# Avoid missing fonts and inherited CJK line grid stretching.
for p in D.paragraphs:
 for r in p.runs:
  rp=r._r.get_or_add_rPr();fonts=rp.get_or_add_rFonts();fonts.set(qn('w:eastAsia'),'黑体' if p.style.name in ['Title','Heading 1','Heading 2','Heading 3'] else ('仿宋' if p.style.name in ['Body Text','作者'] else '宋体'))
  for a in list(fonts.attrib):
   if 'Theme' in a:del fonts.attrib[a]
# Source grid is retained, paragraphs opt out so the exact font size controls line spacing.
for p in D._element.xpath('//w:p'):
 pp=p.find(qn('w:pPr'))
 if pp is None:pp=OxmlElement('w:pPr');p.insert(0,pp)
 snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
settings=D.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
D.save(OUT/'实验三_MySQL读写性能测试报告.docx')
print(OUT/'实验三_MySQL读写性能测试报告.docx')
