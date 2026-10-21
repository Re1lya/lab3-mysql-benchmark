"""Build the submission report from the 18 selected, verified cases."""
import json,importlib.util
from pathlib import Path
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
EV=ROOT/'results/实验数据'
if not EV.exists():EV=ROOT/'output/playwright/lab3-final-20261018'
spec=importlib.util.spec_from_file_location('layout',Path(__file__).with_name('report-layout.py'))
layout=importlib.util.module_from_spec(spec);spec.loader.exec_module(layout);layout.DRAFT=False
layout.doc.settings.odd_and_even_pages_header_footer=False
doc=layout.doc;para=layout.para;heading=layout.heading;sub=layout.sub;page=layout.page;table=layout.table;picture=layout.picture
d=json.loads((EV/'统计与图表/统计结果.json').read_text(encoding='utf8'));rows=d['rows']
assert d['groups']==18 and d['samples']==31200 and all(r['errors']==0 for r in rows)
s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
doc.settings.odd_and_even_pages_header_footer=False
s.top_margin=Inches(.7);s.bottom_margin=Inches(.65);s.left_margin=Inches(.9);s.right_margin=Inches(.9)
s.header_distance=Inches(.28);s.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体')
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(23);doc.styles['Subtitle'].font.size=Pt(13)
doc.styles['Subtitle'].font.italic=False
doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 2'].font.size=Pt(12)
hp=s.header.paragraphs[0];hp.text='Spring 框架课程实验三  |  MySQL 读写性能测试';hp.runs[0].font.size=Pt(8)
fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.add_run('第 ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f);fp.add_run(' 页')
for r in fp.runs:r.font.size=Pt(9)


def results_table(op):
    rs=sorted([r for r in rows if r['operation']==op],key=lambda r:(r['threads'],r['profile']))
    table(['线程数','实现','请求数','均值 ms','P95 ms','≤40 ms比例','达到80%'],[[r['threads'],'JPA' if r['profile']=='jpa' else 'MyBatis',r['samples'],f'{r["mean_ms"]:.2f}',f'{r["p95_ms"]:.2f}',f'{r["within40_pct"]:.2f}%','是' if r['within40_pct']>=80 else '否'] for r in rs],[.65,.8,.7,.85,.85,1.1,.75])

p=doc.add_paragraph('MySQL 读写性能测试',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Spring Data JPA 与 MyBatis',style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para('实验日期：2026年10月18日')
heading('摘要')
para('本实验使用JMeter测试Spring应用的商品读取与新建接口，完成JPA性能测试、MyBatis读写效率比较及两种实现的对比。负载参数参考实验报告样例，读取设5档线程数、写入设4档，两种实现共18组、31200次请求，业务断言均通过。使用Prometheus记录JVM活跃线程、CPU和内存使用率。')
para('关键词：MySQL；Spring Data JPA；MyBatis；JMeter')
heading('1 问题描述')
para('在相同应用、数据库和接口下切换JPA与MyBatis，记录读写请求的响应时间，比较不同线程档位下的测试结果。材料包括原始JTL、JMeter报告、监控数据和过程截图，供验收时逐项核对。')
heading('2 实验分析')
sub('2.1 实验环境')
table(['服务器','用途','主要软件'],[['CDWY3','编译与压力测试','JDK 21、Maven 3.9.16、JMeter 5.6.3'],['CDWY1','应用服务','Spring Boot 4.1.1、JPA、MyBatis 3.5.19'],['CDWY2','数据库','MySQL 9.2.0'],['CDWY4','监控服务','Prometheus 3.7.3']],[.8,1.5,3.4])
para('四台服务器均为Ubuntu 18.04.6、2核CPU、2 GiB内存，通过私网连接。应用使用Docker Swarm部署，各节点运行node-exporter。初始数据库含商品3912行、关联30660行、上架3914行。每组测试前恢复相同初始数据。')
para('应用Tomcat最大/最小线程300/100，Druid最大/最小/初始连接300/100/100。应用堆256–512 MiB，JMeter堆256–768 MiB，线程栈512 KiB。JPA和MyBatis使用同一应用镜像，通过profile切换；两套实现各16项功能测试通过。')


page('2.2 实验方法与JPA结果')
para('读取接口为GET /shops/{shopId}/products/{id}，商品ID来自CSV；写入接口为POST /shops/1/products，字段按JMX计划生成。JMeter以非GUI模式执行，读取断言HTTP 200且errno=0，写入断言HTTP 201；保存原始JTL并生成HTML报告。')
table(['操作','线程数','升压','循环','两实现合计'],[['读取','1000、1050、1100、1150、1200','10秒','2次','10组22000请求'],['写入','1000、1100、1200、1300','10秒','1次','8组9200请求']],[.55,2.75,.6,.6,1.2])
para('本报告统一选用原采集第一轮，共18组31200次请求。例如1000-10-2表示1000线程在10秒内启动，每线程执行2次请求。参考样例，以至少80%的请求成功且响应时间不超过40 ms作为比较条件。')
para('后续按参数逐档展示MyBatis的JMeter网页截图：结果概览、Response Times Over Time和Response Time Percentiles。趋势和百分位均采用JMeter记录的完整响应时间elapsed。PASS表示业务断言通过；40 ms内成功比例另由原始JTL计算。')
sub('JPA性能测试结果')
jpa=sorted([r for r in rows if r['profile']=='jpa'],key=lambda r:(r['operation'],r['threads']))
table(['操作','线程数','请求数','平均响应 ms','≤40 ms比例'],[['读取' if r['operation']=='ReadProduct' else '写入',r['threads'],r['samples'],f'{r["mean_ms"]:.2f}',f'{r["within40_pct"]:.2f}%'] for r in jpa],[.65,.8,.85,1.5,1.4])
para('JPA的9组断言均通过。5组读取和1000、1100、1200线程写入达到响应时间比较条件；1300线程写入平均响应为3232.80 ms，40 ms内成功比例为0%。各组原始JTL和JMeter图表均已保存。')

figure=0
def screenshot_case(r,section):
    global figure
    op='读取' if r['operation']=='ReadProduct' else '写入'
    page(section)
    sub(f'{r["threads"]}-10-{r["loops"]}')
    para(f'MyBatis{op}，{r["threads"]}线程、10秒升压、{r["loops"]}次循环，共{r["samples"]}次请求，断言失败0次。平均响应时间为{r["mean_ms"]:.2f} ms；40 ms内成功比例为{r["within40_pct"]:.2f}%，'+('达到' if r['within40_pct']>=80 else '未达到')+'80%的比较条件。')
    for suffix in ['结果概览','响应时间趋势','响应时间百分位']:
        path=EV/'过程截图/原生图表'/(r['label']+'_'+suffix+'.png')
        p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
        p.add_run().add_picture(str(path),width=Inches(6.4))
    figure+=1
    p=para(f'图{figure}  MyBatis{op} {r["threads"]}-10-{r["loops"]} 的JMeter原生测试结果。')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:run.font.size=Pt(9)
for op,title in [('ReadProduct','2.3 ReadProduct API测试'),('WriteProduct','2.4 WriteProduct API测试')]:
    for r in sorted([x for x in rows if x['profile']=='mybatis' and x['operation']==op],key=lambda x:x['threads']):
        screenshot_case(r,title)

page('2.5 读写效率与框架对比')
sub('MyBatis读写比较')
table(['线程数','读取均值 ms','写入均值 ms','读取/写入循环'],[[n,f'{next(r["mean_ms"] for r in rows if r["profile"]=="mybatis" and r["operation"]=="ReadProduct" and r["threads"]==n):.2f}',f'{next(r["mean_ms"] for r in rows if r["profile"]=="mybatis" and r["operation"]=="WriteProduct" and r["threads"]==n):.2f}','2次 / 1次'] for n in [1000,1100,1200]],[.8,1.5,1.5,1.8])
para('在读写共有的1000、1100、1200线程档位下，MyBatis读取平均耗时均小于写入。读取与写入的循环次数按样例分别设置为2次和1次。1300线程写入平均响应为102.03 ms，40 ms内成功比例为71.77%。')
sub('JPA与MyBatis比较')
comp=[]
for r in jpa:
    m=next(x for x in rows if x['profile']=='mybatis' and x['operation']==r['operation'] and x['threads']==r['threads'])
    comp.append(['读取' if r['operation']=='ReadProduct' else '写入',r['threads'],f'{r["mean_ms"]:.2f}',f'{m["mean_ms"]:.2f}'])
table(['操作','线程数','JPA均值 ms','MyBatis均值 ms'],comp,[.7,.9,1.8,1.8])
para('同一操作和线程档位下，两种实现使用相同的升压时间、循环次数、请求生成规则和初始数据，仅切换持久层。在本次各档读取和写入记录中，MyBatis平均响应时间均低于JPA。')
page('2.6 运行监控')
para('以下为MyBatis 1300-10-1写入期间的Prometheus历史监控页面，分别显示JVM活跃线程、应用节点CPU和内存使用率。截图于2026年10月21日按原测试时间补截，测量数据仍来自10月18日该轮实验。页面时间为UTC，北京时间需加8小时。')
for title in ['活跃线程','CPU使用率','内存使用率']:
    if title=='内存使用率':page('2.6 运行监控与实验总结')
    sub(title)
    p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
    p.add_run().add_picture(str(EV/'过程截图/运行监控'/('MyBatis写入_1300线程_第1轮_'+title+'.png')),width=Inches(6.4))
para('CPU使用率采用rate[5m]规则；内存使用率为[1−(MemFree+Buffers+Cached)/MemTotal]×100%。抓取间隔2秒，记录规则组每30秒计算一次。三项查询的原始记录位于对应组的监控JSON。')
heading('3 总结')
para('本实验完成JPA读写性能测试、MyBatis读写比较及两种实现的对比。18组共31200次请求，断言失败为0；10组读取及6组写入达到响应时间比较条件，两个1300线程写入组未达到。原始JTL、JMeter报告和运行监控均已保存。')
para('程序地址：https://github.com/Re1lya/lab3-mysql-benchmark')

sub('参考资料')
para('[1] 实验三：MySQL读写性能测试，第1–4页。')
para('[2] 实验报告样例，第1–10页。')
out=ROOT/'report' if (ROOT/'results/实验数据').exists() else ROOT/'output/playwright/lab3-teacher-20261018/deliverables'
out.mkdir(parents=True,exist_ok=True);target=out/'实验三_MySQL读写性能测试报告.docx';doc.save(target);print(target)
