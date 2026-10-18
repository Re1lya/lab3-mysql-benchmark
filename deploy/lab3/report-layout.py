"""Create the experiment report from verified result JSON, with no invented measurements."""
import argparse,json,statistics
from pathlib import Path
from datetime import datetime,timezone,timedelta
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
EV=ROOT/'output/playwright/lab3-teacher-20261018'
REPO='https://github.com/Re1lya/lab3-mysql-benchmark'
TZ=timezone(timedelta(hours=8))
doc=Document()

def para(text='',style=None):
    p=doc.add_paragraph(text,style)
    return p

def heading(text):doc.add_heading(text,level=1)
def sub(text):doc.add_heading(text,level=2)
def page(title):doc.add_page_break();heading(title)
def fmt(v,d=1):return f'{v:,.{d}f}'

def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [6.7/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
    for row in rows:
        cells=t.add_row().cells
        for c,text in zip(cells,row):c.text=str(text)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if i==0:trPr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd')
            sh.set(qn('w:fill'),'E5EBF1' if i==0 else ('F7F9FB' if i%2==0 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');margins.append(x)
            pr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.08
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9.5);r.font.bold=i==0
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    para().paragraph_format.space_after=Pt(1)
    return t

def picture(path,caption,width=6.6):
    path=Path(path)
    if not path.exists():
        if not DRAFT:raise FileNotFoundError(path)
        para('草稿阶段尚未生成本图。');return
    p=para();p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Inches(width))
    p=para(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(9)
    for r in p.runs:r.font.size=Pt(9)
